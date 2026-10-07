import io
import hashlib
from pathlib import Path
import pytest
from gulms.slide_extractor import SlideExtractor, is_template_asset, is_template_text, TEMPLATE_CORPUS_PREFIXES


def test_template_corpus_hashes_registry():
    """Verify pre-computed university branding hashes are correctly registered."""
    assert len(TEMPLATE_CORPUS_PREFIXES) >= 4
    # G-SCALE badge hash prefix
    assert "305ff7cbda57d943" in TEMPLATE_CORPUS_PREFIXES
    # An asset with this OCR phrase is also recognized
    assert is_template_asset(b"test_bytes", ocr_text="Galgotias University Active Learning Ecosystem")


def test_template_text_filter():
    """Verify boilerplate chrome text patterns are accurately filtered."""
    assert is_template_text("Galgotias University")
    assert is_template_text("G-SCALE")
    assert is_template_text("Student-Centred Active Learning Ecosystem")
    assert is_template_text("Think-Pair-Share")
    assert is_template_text("[5-mins]")
    assert is_template_text("23")  # slide number
    # Real academic content must NOT be filtered
    assert not is_template_text("Database System Vs File System")
    assert not is_template_text("Flowchart for the for loop")
    assert not is_template_text("Addr(A[Index]) = B + W * (Index - LB)")


def test_slide_extractor_standalone_synthetic():
    """Verify end-to-end slide extraction and table translation on a standalone synthesized presentation."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        pytest.skip("python-pptx not available in current test environment")

    prs = Presentation()
    # Slide 1: Title slide
    title_slide_layout = prs.slide_layouts[0]
    s1 = prs.slides.add_slide(title_slide_layout)
    s1.shapes.title.text = "Introduction to Relational Databases"
    s1.placeholders[1].text = "Core concepts and ACID properties"

    # Slide 2: Table slide
    blank_layout = prs.slide_layouts[6]
    s2 = prs.slides.add_slide(blank_layout)
    # Add a title text box
    tb = s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tb.text = "Database Comparison"
    # Add a 3x2 table (rows=3, cols=2)
    tbl_shape = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    table = tbl_shape.table
    table.cell(0, 0).text = "Paradigm"
    table.cell(0, 1).text = "ACID Support"
    table.cell(1, 0).text = "RDBMS"
    table.cell(1, 1).text = "Strict"
    table.cell(2, 0).text = "NoSQL"
    table.cell(2, 1).text = "Eventual"

    bio = io.BytesIO()
    prs.save(bio)
    deck_bytes = bio.getvalue()

    extractor = SlideExtractor()
    slides = extractor.extract_deck(deck_bytes)

    assert len(slides) == 2
    assert slides[0]["slide_num"] == 1
    assert slides[0]["title"] == "Introduction to Relational Databases"
    assert "ACID properties" in slides[0]["markdown"]

    assert slides[1]["slide_num"] == 2
    assert slides[1]["title"] == "Database Comparison"
    assert len(slides[1]["tables"]) == 1
    assert "| Paradigm | ACID Support |" in slides[1]["markdown"]
    assert "| RDBMS | Strict |" in slides[1]["markdown"]
    assert "| NoSQL | Eventual |" in slides[1]["markdown"]


def test_drawingml_connector_routing(benchmark_dir):
    """Verify DrawingML connectors produce directed Mermaid graphs with condition labels."""
    jp_deck = benchmark_dir / "JP_Lecture___04_If_switch_loops_p.pptx"
    if not jp_deck.exists():
        pytest.skip(f"Benchmark deck not found at {jp_deck}. Set BENCHMARK_DIR to run benchmark tests.")
    
    deck_bytes = jp_deck.read_bytes()
    extractor = SlideExtractor()
    slides = extractor.extract_deck(deck_bytes)
    
    # Slide 23 is the for-loop flowchart
    s23 = slides[22]
    assert s23["slide_num"] == 23
    assert s23["title"] == "Flowchart for the for loop"
    assert len(s23["graph_nodes"]) == 4
    assert len(s23["graph_edges"]) >= 3
    
    # Check that condition labels are preserved
    edge_labels = [lbl for _, _, lbl in s23["graph_edges"] if lbl]
    assert "true" in edge_labels
    
    # Check Mermaid syntax
    assert "```mermaid\nflowchart TD" in s23["markdown"]
    assert 'node_184325{"condition?"}' in s23["markdown"]
    assert " -- true --> " in s23["markdown"]


def test_native_table_extraction(benchmark_dir):
    """Verify native <a:tbl> tables convert to pipe-delimited Markdown tables."""
    ds_deck = benchmark_dir / "DS_Lec_05_DS_pptx.pptx"
    if not ds_deck.exists():
        pytest.skip(f"Benchmark deck not found at {ds_deck}. Set BENCHMARK_DIR to run benchmark tests.")
        
    deck_bytes = ds_deck.read_bytes()
    extractor = SlideExtractor()
    slides = extractor.extract_deck(deck_bytes)
    
    # Slide 16 is the Summary table
    s16 = slides[15]
    assert s16["slide_num"] == 16
    assert s16["title"] == "Summary"
    assert len(s16["tables"]) == 1
    
    tbl = s16["tables"][0]
    assert len(tbl) == 6  # 1 header + 5 rows
    assert len(tbl[0]) == 4  # 4 columns
    assert "Array Type" in tbl[0]
    assert "Formula" in tbl[0]
    
    # Check markdown table structure
    assert "| Array Type | Order | Formula | Parameters |" in s16["markdown"]
    assert "| --- | --- | --- | --- |" in s16["markdown"]
    assert "1D Array" in s16["markdown"]
    assert "3D Array" in s16["markdown"]


def test_group_shape_recursive_unrolling(benchmark_dir):
    """Verify shapes nested in <p:grpSp> are unrolled and accessible."""
    jp_deck = benchmark_dir / "JP_Lecture___04_If_switch_loops_p.pptx"
    if not jp_deck.exists():
        pytest.skip(f"Benchmark deck not found at {jp_deck}. Set BENCHMARK_DIR to run benchmark tests.")
        
    from pptx import Presentation
    prs = Presentation(str(jp_deck))
    slide = prs.slides[22]
    
    extractor = SlideExtractor()
    flat = extractor._flatten_shapes(slide.shapes)
    # Slide 23 has 5 top-level shapes, but inside Group 3 there are 13 sub-shapes
    assert len(flat) > 10
