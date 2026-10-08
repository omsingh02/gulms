"""Regenerates the synthetic decks used by the extractor tests.

Needs `python-pptx` and `Pillow`. The expected-output JSON files were produced by the
original Python extractor (see git history, `gulms/slide_extractor.py`) and are checked
in; the Rust tests compare against them. Only rerun this if you intentionally change a fixture.
"""
import io
import sys
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from PIL import Image, ImageDraw


def png(width, height, color):
    img = Image.new("RGB", (width, height), color)
    ImageDraw.Draw(img).rectangle([5, 5, width - 6, height - 6], outline=(0, 0, 0), width=3)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def simple_deck():
    """The two-slide deck from the original Python test-suite."""
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Introduction to Relational Databases"
    s1.placeholders[1].text = "Core concepts and ACID properties"

    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tb.text = "Database Comparison"
    table = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate([("Paradigm", "ACID Support"), ("RDBMS", "Strict"), ("NoSQL", "Eventual")]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    return prs


def rich_deck():
    prs = Presentation()

    # 1. Title + multi-level bullets with a soft line break and template chrome text
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Normal   Forms"
    body = s.placeholders[1].text_frame
    body.text = "First normal form"
    p = body.add_paragraph(); p.text = "Atomic values"; p.level = 1
    p = body.add_paragraph(); p.text = "No repeating groups"; p.level = 2
    p = body.add_paragraph(); p.text = "Second line one"; p.add_line_break(); p.add_run().text = "second line two"
    p = body.add_paragraph(); p.text = "Galgotias University"
    p = body.add_paragraph(); p.text = "17"

    # 2. Title-only slide, table with an empty row and ragged text, small icon + real figure
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Keys"
    t = s.shapes.add_table(4, 3, Inches(0.5), Inches(1.8), Inches(8), Inches(2)).table
    rows = [("Key", "Meaning", "Example"), ("PK", "Unique row id", "roll_no"), ("", "", ""), ("FK", "Link\nto parent", "dept_id")]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    s.shapes.add_picture(png(240, 140, (200, 220, 255)), Inches(0.5), Inches(4.2), Inches(2.4), Inches(1.4))
    s.shapes.add_picture(png(40, 40, (255, 0, 0)), Inches(8), Inches(0.2))

    # 3. Flowchart: decision + two processes joined by connectors, a branch label, a group
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(0.8)); tb.text = "Flowchart for the loop"
    start = s.shapes.add_shape(MSO_SHAPE.FLOWCHART_PROCESS, Inches(1), Inches(1.5), Inches(2), Inches(0.8)); start.text = "i = 0"
    cond = s.shapes.add_shape(MSO_SHAPE.FLOWCHART_DECISION, Inches(1), Inches(3), Inches(2), Inches(1.2)); cond.text = "i < 10 ?"
    body_ = s.shapes.add_shape(MSO_SHAPE.FLOWCHART_PROCESS, Inches(5), Inches(3.1), Inches(2), Inches(0.8)); body_.text = "print(i)"
    c1 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2), Inches(2.3), Inches(2), Inches(3))
    c2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(3.6), Inches(5), Inches(3.5))
    lbl = s.shapes.add_textbox(Inches(3.6), Inches(3.0), Inches(0.6), Inches(0.4)); lbl.text = "yes"
    grp = s.shapes.add_group_shape()
    note = grp.shapes.add_textbox(Inches(5), Inches(5), Inches(3), Inches(0.6)); note.text = "Loop runs ten times"

    # 4. Boxes with no arrows (the original reported these as a flowchart anyway)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(0.8)); tb.text = "Two plain boxes"
    a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1)); a.text = "Left box"
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(2), Inches(3), Inches(1)); b.text = "Right box"
    return prs


def save(prs, path):
    prs.save(path)
    print("wrote", path, file=sys.stderr)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    save(simple_deck(), f"{out}/simple.pptx")
    save(rich_deck(), f"{out}/rich.pptx")
